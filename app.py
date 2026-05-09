"""MarineRAG-DSS – Streamlit web application.

Provides a conversational chatbot interface that uses a RAG pipeline
to answer questions about marine resource management, drawing from
environmental data, marine laws, and fisheries knowledge.
"""

from __future__ import annotations

import csv
import io
import logging
import os
import re
import smtplib
from datetime import datetime, time, timezone
from email.message import EmailMessage
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlencode, urlparse, urlunparse

import streamlit as st
from bson import ObjectId

from src.auth import fetch_user_info, generate_state, get_login_url, parse_state
from src.chatbot import MarineChatbot
from src.config import LLMMode, config
from src.db import (
    append_message,
    authenticate_email_user,
    create_chat_session,
    create_email_user,
    create_email_verification_token,
    create_password_reset_token,
    ensure_indexes,
    get_admin_stats,
    get_admin_analytics,
    get_session_messages,
    get_user_by_email,
    list_sessions,
    list_sessions_filtered,
    list_messages_filtered,
    list_user_ids_by_role,
    list_users,
    reset_password_with_token,
    update_user_role,
    upsert_user,
    verify_email_token,
)
from src.permissions import NON_ADMIN_ROLES, ROLE_LABELS, ROLES, get_permissions, normalize_role
from src.video_links import load_guide_video_links

logging.basicConfig(
    level=getattr(logging, config.LOG_LEVEL, logging.INFO),
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
logger = logging.getLogger(__name__)

_CITATION_PATTERN = re.compile(
    r"\s*\((?:[^()]*?(?:data[\\/]|sheet:|row:|nguon|nguồn|source|https?://|\[[0-9]+\]))[^()]*\)",
    re.IGNORECASE,
)
_GUIDE_QUERY_HINTS = (
    "huong dan",
    "hướng dẫn",
    "su dung",
    "sử dụng",
    "cach dung",
    "cách dùng",
    "how to use",
    "user guide",
    "manual",
)
_YOUTUBE_URL_PATTERN = re.compile(
    r"https?://(?:www\.)?(?:youtube\.com|youtu\.be)/\S+",
    re.IGNORECASE,
)
_EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _format_role_label(value: str | None) -> str:
    if not value:
        return ""
    return str(ROLE_LABELS.get(value, value))

# ---------------------------------------------------------------------------
# Page configuration
# ---------------------------------------------------------------------------
st.set_page_config(
    page_title=config.APP_TITLE,
    page_icon="🌊",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ---------------------------------------------------------------------------
# Session state helpers
# ---------------------------------------------------------------------------

def _init_session_state() -> None:
    if "messages" not in st.session_state:
        st.session_state.messages = []
    if "chatbot" not in st.session_state:
        st.session_state.chatbot = None
    if "llm_mode" not in st.session_state:
        st.session_state.llm_mode = LLMMode.API
    if "setup_done" not in st.session_state:
        st.session_state.setup_done = False
    if "setup_error" not in st.session_state:
        st.session_state.setup_error = None
    if "user" not in st.session_state:
        st.session_state.user = None
    if "active_session_id" not in st.session_state:
        st.session_state.active_session_id = None
    if "oauth_state" not in st.session_state:
        st.session_state.oauth_state = ""
    if "oauth_states" not in st.session_state:
        st.session_state.oauth_states = {}
    if "oauth_provider" not in st.session_state:
        st.session_state.oauth_provider = ""
    if "pending_reset_token" not in st.session_state:
        st.session_state.pending_reset_token = ""
    if "page" not in st.session_state:
        st.session_state.page = "home"
    if "auth_notice" not in st.session_state:
        st.session_state.auth_notice = None


def _get_query_params() -> dict:
    if hasattr(st, "query_params"):
        return dict(st.query_params)
    get_params = getattr(st, "experimental_get_query_params", None)
    if callable(get_params):
        params = get_params()
        if isinstance(params, dict):
            return params
    return {}


def _clear_query_params() -> None:
    if hasattr(st, "query_params"):
        st.query_params.clear()
        return
    set_params = getattr(st, "experimental_set_query_params", None)
    if callable(set_params):
        set_params()


def _smtp_ready() -> bool:
    return bool(config.SMTP_HOST and config.SMTP_FROM)


def _send_email(to_address: str, subject: str, body: str) -> bool:
    if not _smtp_ready():
        return False
    message = EmailMessage()
    message["From"] = config.SMTP_FROM
    message["To"] = to_address
    message["Subject"] = subject
    message.set_content(body)

    try:
        if config.SMTP_USE_TLS:
            with smtplib.SMTP(config.SMTP_HOST, config.SMTP_PORT, timeout=10) as smtp:
                smtp.starttls()
                if config.SMTP_USER:
                    smtp.login(config.SMTP_USER, config.SMTP_PASSWORD)
                smtp.send_message(message)
        else:
            with smtplib.SMTP_SSL(config.SMTP_HOST, config.SMTP_PORT, timeout=10) as smtp:
                if config.SMTP_USER:
                    smtp.login(config.SMTP_USER, config.SMTP_PASSWORD)
                smtp.send_message(message)
    except Exception as exc:
        logger.warning("Failed to send email to %s: %s", to_address, exc)
        return False
    return True


def _build_app_link(param: str, token: str) -> str:
    base_url = (config.APP_BASE_URL or config.OAUTH_REDIRECT_URI).strip()
    if not base_url:
        base_url = "http://localhost:8501"
    parsed = urlparse(base_url)
    query = parse_qs(parsed.query)
    query[param] = [token]
    new_query = urlencode(query, doseq=True)
    return urlunparse(parsed._replace(query=new_query))


def _set_page(page: str) -> None:
        st.session_state.page = page


def _set_auth_notice(level: str, message: str) -> None:
        st.session_state.auth_notice = (level, message)


def _render_auth_notice() -> None:
        notice = st.session_state.auth_notice
        if not notice:
                return
        level, message = notice
        if level == "success":
                st.success(message)
        elif level == "error":
                st.error(message)
        else:
                st.info(message)
        st.session_state.auth_notice = None


def _inject_public_styles() -> None:
        st.markdown(
                """
<style>
@import url('https://fonts.googleapis.com/css2?family=Playfair+Display:wght@600;700&family=Space+Grotesk:wght@400;600;700&display=swap');

:root {
    --ocean-0: #e0f7fa;
    --ocean-1: #b2ebf2;
    --ocean-2: #80deea;
    --sea: #00acc1;
    --sea-2: #00838f;
    --sand: #005b9f;
    --fog: #1a365d;
    --glass: rgba(255, 255, 255, 0.6);
}

.stApp {
    background: radial-gradient(1200px 600px at 10% -10%, rgba(128, 222, 234, 0.5), transparent 60%),
                radial-gradient(1000px 700px at 90% 0%, rgba(0, 172, 193, 0.3), transparent 55%),
                linear-gradient(160deg, var(--ocean-0), var(--ocean-1));
    color: var(--fog);
}

.block-container {
    padding-top: 2.2rem;
}

h1, h2, h3, h4, h5, h6 {
    font-family: 'Playfair Display', serif;
    color: var(--sand);
}

p, li, span, div {
    font-family: 'Space Grotesk', sans-serif;
}

.public-hero {
    padding: 2rem 2.2rem;
    border-radius: 24px;
    background: linear-gradient(130deg, rgba(255, 255, 255, 0.8), rgba(224, 247, 250, 0.9));
    border: 1px solid rgba(0, 172, 193, 0.2);
    box-shadow: 0 12px 40px rgba(0, 91, 159, 0.1);
}

.hero-kicker {
    letter-spacing: 0.22em;
    text-transform: uppercase;
    font-size: 0.75rem;
    color: var(--sea-2);
    margin-bottom: 0.75rem;
}

.hero-title {
    font-size: 2.6rem;
    margin-bottom: 0.8rem;
}

.hero-sub {
    font-size: 1.05rem;
    line-height: 1.7;
    color: var(--fog);
}

.glass-card {
    padding: 1.4rem 1.2rem;
    border-radius: 18px;
    background: var(--glass);
    border: 1px solid rgba(255, 255, 255, 0.8);
    box-shadow: 0 8px 32px rgba(0, 91, 159, 0.08);
}

.nav-tag {
    text-align: right;
    font-size: 0.85rem;
    color: var(--sea-2);
    padding-top: 0.6rem;
    font-weight: 600;
}

.page-title {
    font-size: 2rem;
    margin-bottom: 0.6rem;
}

.fade-in {
    animation: fadeUp 0.8s ease both;
}

@keyframes fadeUp {
    from { opacity: 0; transform: translateY(12px); }
    to { opacity: 1; transform: translateY(0); }
}

div.stButton > button, div.stFormSubmitButton > button {
    border-radius: 999px !important;
    border: 1px solid var(--sea) !important;
    background-color: white !important;
    color: var(--sea-2) !important;
    padding: 0.45rem 1.2rem !important;
    font-weight: 600 !important;
    transition: all 0.2s ease !important;
}

div.stButton > button:hover, div.stFormSubmitButton > button:hover {
    background-color: var(--sea) !important;
    color: white !important;
    border-color: var(--sea) !important;
    box-shadow: 0 4px 12px rgba(0, 172, 193, 0.3) !important;
}

/* Override Streamlit text inputs for light theme */
div[data-baseweb="input"] > div, div[data-baseweb="baseInput"], div[data-baseweb="select"] > div {
    background-color: white !important;
    border: 1px solid rgba(0, 172, 193, 0.4) !important;
    border-radius: 8px !important;
}
div[data-baseweb="input"] input, div[data-baseweb="select"] span {
    color: var(--fog) !important;
    -webkit-text-fill-color: var(--fog) !important;
}
div[data-baseweb="input"] input::placeholder {
    color: #6482a5 !important;
    opacity: 1 !important;
}
div[data-baseweb="input"] > div:focus-within, div[data-baseweb="select"] > div:focus-within {
    border-color: var(--sea-2) !important;
    box-shadow: 0 0 0 1px var(--sea-2) !important;
}
label, p[data-testid="stMarkdownContainer"] {
    color: var(--fog) !important;
}
.st-emotion-cache-16idsys p {
    color: var(--fog) !important;
}

/* Hide Streamlit top header and decoration */
header[data-testid="stHeader"] {
    display: none !important;
}
div[data-testid="stDecoration"] {
    display: none !important;
}
</style>
""",
                unsafe_allow_html=True,
        )


def _render_top_nav() -> None:
        nav_cols = st.columns([1.3, 1.4, 1.4, 6])
        with nav_cols[0]:
                if st.button("Trang chủ", key="nav_home"):
                        _set_page("home")
                        st.rerun()
        with nav_cols[1]:
                if st.button("Đăng nhập", key="nav_login"):
                        _set_page("login")
                        st.rerun()
        with nav_cols[2]:
                if st.button("Đăng ký", key="nav_register"):
                        _set_page("register")
                        st.rerun()
        with nav_cols[3]:
                st.markdown("<div class='nav-tag'>MarineRAG DSS</div>", unsafe_allow_html=True)


def _handle_auth_query_params() -> None:
        params = _get_query_params()
        verify_token = params.get("verify", "")
        reset_token = params.get("reset", "")
        if isinstance(verify_token, list):
                verify_token = verify_token[0] if verify_token else ""
        if isinstance(reset_token, list):
                reset_token = reset_token[0] if reset_token else ""

        if verify_token:
                if verify_email_token(verify_token):
                        _set_auth_notice("success", "Email verified. You can sign in now.")
                else:
                        _set_auth_notice("error", "Verification link is invalid or expired.")
                _set_page("login")
                _clear_query_params()

        if reset_token:
                st.session_state.pending_reset_token = reset_token
                _set_page("forgot")
                _clear_query_params()


def _strip_inline_citations(text: str) -> str:
    if not text:
        return text
    cleaned = _CITATION_PATTERN.sub("", text)
    cleaned = re.sub(r"[ \t]+\n", "\n", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def _strip_youtube_urls(text: str) -> str:
    if not text:
        return text
    cleaned = _YOUTUBE_URL_PATTERN.sub("", text)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def _is_guide_query(text: str) -> bool:
    lowered = (text or "").strip().lower()
    return any(hint in lowered for hint in _GUIDE_QUERY_HINTS)


def _guide_video_urls_for_question(question: str) -> list[str]:
    if not _is_guide_query(question):
        return []

    links = load_guide_video_links(config.GUIDE_VIDEO_LINKS_PATH)
    urls: list[str] = []
    seen: set[str] = set()
    for item in links:
        url = str(item.get("url", "")).strip()
        if not url or url in seen:
            continue
        seen.add(url)
        urls.append(url)
    return urls


def _extract_youtube_urls(text: str) -> list[str]:
    if not text:
        return []

    urls: list[str] = []
    seen: set[str] = set()
    for match in _YOUTUBE_URL_PATTERN.findall(text):
        url = match.strip().rstrip(").,;]")
        if url in seen:
            continue
        seen.add(url)
        urls.append(url)
    return urls


def _render_embedded_videos(text: str, question: str | None = None) -> None:
    urls = _extract_youtube_urls(text)
    for url in _guide_video_urls_for_question(question or ""):
        if url not in urls:
            urls.append(url)
    if not urls:
        return

    for idx, url in enumerate(urls, start=1):
        st.video(url)


def _setup_chatbot(llm_mode: str, force_rebuild: bool = False) -> None:
    """Instantiate and set up the MarineChatbot; store in session state."""
    with st.spinner("Loading knowledge base and preparing KG2RAG pipeline…"):
        try:
            bot = MarineChatbot(
                llm_mode=llm_mode,
                force_rebuild=force_rebuild,
            )
            bot.setup()
            st.session_state.chatbot = bot
            st.session_state.llm_mode = llm_mode
            st.session_state.setup_done = True
            st.session_state.setup_error = None
            logger.info("Chatbot setup complete (mode=%s).", llm_mode)
        except Exception as exc:
            logger.exception("Chatbot setup failed.")
            st.session_state.setup_done = False
            st.session_state.setup_error = str(exc)


# ---------------------------------------------------------------------------
# Sidebar
# ---------------------------------------------------------------------------

def _render_sidebar() -> None:
    st.sidebar.title("⚙️ Configuration")

    user = st.session_state.user
    if user:
        user_role = normalize_role(user.get("role", ""))
        role_label = ROLE_LABELS.get(user_role, user.get("role", ""))
        st.sidebar.markdown(f"**User:** {user.get('name','')} ({user.get('email','')})")
        st.sidebar.markdown(f"**Role:** {role_label}")
        if st.sidebar.button("Log out", use_container_width=True):
            st.session_state.user = None
            st.session_state.active_session_id = None
            st.session_state.messages = []
            st.session_state.page = "home"
            st.rerun()

    st.sidebar.markdown("### LLM Mode")
    llm_mode = st.sidebar.radio(
        "Select LLM backend:",
        options=[LLMMode.API, LLMMode.LOCAL],
        format_func=lambda m: "☁️ Gemini API" if m == LLMMode.API else "💻 Local (Ollama)",
        index=0 if st.session_state.llm_mode == LLMMode.API else 1,
        key="llm_mode_radio",
    )

    if llm_mode == LLMMode.API:
        api_key = st.sidebar.text_input(
            "Gemini API Key",
            value=os.getenv("GEMINI_API_KEY", ""),
            type="password",
            help="Required for API mode. Set via GEMINI_API_KEY env variable or enter here.",
        )
        if api_key:
            os.environ["GEMINI_API_KEY"] = api_key

        st.sidebar.markdown(
            f"**Model:** `{config.GEMINI_MODEL}`  \n"
            f"**Embeddings:** `{config.GEMINI_EMBEDDING_MODEL}`"
        )
    else:
        embeddings_provider = (config.LOCAL_EMBEDDINGS_PROVIDER or "ollama").strip().lower()
        if embeddings_provider in {"hf", "huggingface", "sentence-transformers", "sbert"}:
            embeddings_info = f"**Embeddings (HF):** `{config.HF_EMBEDDING_MODEL}`"
            embeddings_note = "_Embeddings run locally via sentence-transformers._"
        else:
            embeddings_info = f"**Embeddings (Ollama):** `{config.OLLAMA_EMBEDDING_MODEL}`"
            embeddings_note = ""

        ollama_note = "_Ensure Ollama is running locally with the required models._"
        notes = f"{embeddings_note}  \n{ollama_note}" if embeddings_note else ollama_note
        st.sidebar.markdown(
            f"**Ollama URL:** `{config.OLLAMA_BASE_URL}`  \n"
            f"**Model:** `{config.OLLAMA_MODEL}`  \n"
            f"{embeddings_info}  \n\n"
            f"{notes}"
        )

    st.sidebar.divider()

    st.sidebar.markdown("### Knowledge Base")
    force_rebuild = st.sidebar.checkbox(
        "Rebuild vector index",
        value=False,
        help="Force re-indexing of all documents (slower; use after adding new data).",
    )

    if st.sidebar.button("🚀 Initialise / Reinitialise", use_container_width=True):
        _setup_chatbot(llm_mode, force_rebuild=force_rebuild)

    if st.session_state.setup_done:
        st.sidebar.success("✅ Chatbot ready")
    elif st.session_state.setup_error:
        st.sidebar.error(f"❌ Setup failed:\n{st.session_state.setup_error}")
    else:
        st.sidebar.info("Click **Initialise** to start.")

    st.sidebar.divider()
    st.sidebar.markdown("### Session")
    if st.sidebar.button("🗑️ Clear chat history", use_container_width=True):
        st.session_state.messages = []
        if st.session_state.chatbot:
            st.session_state.chatbot.reset_history()
        st.rerun()

    st.sidebar.divider()
    st.sidebar.markdown(
        "**Data sources:**\n"
        "- 🌊 Environmental data\n"
        "- ⚖️ Marine laws & regulations\n"
        "- 🐟 Fisheries & aquaculture data\n\n"
        "*Add more `.txt`, `.pdf`, or `.xls/.xlsx` files to `data/sample_docs/` "
        "and reinitialise.*"
    )


# ---------------------------------------------------------------------------
# Main page
# ---------------------------------------------------------------------------

def _render_header() -> None:
    col1, col2 = st.columns([1, 6])
    with col1:
        st.markdown("# 🌊")
    with col2:
        st.title(config.APP_TITLE)
        st.caption(
            "A knowledge-graph-guided Retrieval-Augmented Generation (KG2RAG) "
            "Decision Support System for marine resource management — powered by "
            "multi-source environmental, legal, and fisheries knowledge."
        )
    st.divider()


def _handle_oauth_callback() -> None:
    params = _get_query_params()
    code = params.get("code", "")
    state = params.get("state", "")
    if isinstance(code, list):
        code = code[0] if code else ""
    if isinstance(state, list):
        state = state[0] if state else ""
    if not code or not state:
        return

    provider, token = parse_state(state)
    expected_token = st.session_state.oauth_states.get(provider)
    if expected_token and token != expected_token:
        st.error("Invalid OAuth state. Please try logging in again.")
        _clear_query_params()
        return

    try:
        profile = fetch_user_info(provider, code, config.OAUTH_REDIRECT_URI)
    except Exception as exc:
        st.error(f"OAuth login failed: {exc}")
        _clear_query_params()
        return

    provider_id = profile.get("sub") or profile.get("id") or ""
    email = profile.get("email", "")
    name = profile.get("name") or profile.get("preferred_username") or email
    role = config.DEFAULT_USER_ROLE
    user = upsert_user(provider, provider_id, email, name, role)
    st.session_state.user = user
    _clear_query_params()


def _render_home() -> None:
    _render_top_nav()
    
    st.markdown("<div style='height: 3rem;'></div>", unsafe_allow_html=True)
    
    st.markdown(
        """
<div class="public-hero fade-in" style="text-align: center; max-width: 800px; margin: 0 auto;">
  <div class="hero-kicker">MARINERAG DSS</div>
  <div class="hero-title">Hệ thống Hỗ trợ Ra Quyết định Quản lý Tài nguyên Biển</div>
  <div class="hero-sub">
    Khai thác sức mạnh của trí tuệ nhân tạo (AI) và hệ tri thức (Knowledge Graph) để phân tích dữ liệu môi trường, pháp lý và thủy sản. Trợ lý ảo của chúng tôi giúp bạn trả lời câu hỏi chuyên sâu, trích dẫn nguồn uy tín và hỗ trợ ra quyết định chính xác, nhanh chóng.
  </div>
</div>
""",
        unsafe_allow_html=True,
    )
    
    st.markdown("<div style='height: 2rem;'></div>", unsafe_allow_html=True)

    cta_cols = st.columns([1, 1, 1])
    with cta_cols[1]:
        col1, col2 = st.columns(2)
        with col1:
            if st.button("Bắt đầu ngay", key="home_login", use_container_width=True):
                _set_page("login")
                st.rerun()
        with col2:
            if st.button("Đăng ký tài khoản", key="home_register", use_container_width=True):
                _set_page("register")
                st.rerun()

    st.markdown("<div style='height: 4rem;'></div>", unsafe_allow_html=True)

    st.markdown("<h3 style='text-align: center; margin-bottom: 2rem; color: var(--sand);'>Tính Năng Nổi Bật</h3>", unsafe_allow_html=True)

    feature_cols = st.columns(3)
    with feature_cols[0]:
        st.markdown(
            """
<div class="glass-card fade-in" style="height: 100%;">
  <h3 style="margin-bottom: 1rem;">🧠 RAG & Knowledge Graph</h3>
  <p>Hệ thống tự động kết nối tài liệu rời rạc thành mạng lưới tri thức, giúp chatbot hiểu sâu và đưa ra các câu trả lời toàn diện, logic theo đúng ngữ cảnh hàng hải.</p>
</div>
""",
            unsafe_allow_html=True,
        )
    with feature_cols[1]:
        st.markdown(
            """
<div class="glass-card fade-in" style="height: 100%;">
  <h3 style="margin-bottom: 1rem;">🔍 Minh bạch & Tin cậy</h3>
  <p>Mọi dữ liệu sinh ra đều được đính kèm trích dẫn về nguồn nội bộ, báo cáo môi trường, hoặc văn bản luật pháp cụ thể để dễ dàng kiểm chứng.</p>
</div>
""",
            unsafe_allow_html=True,
        )
    with feature_cols[2]:
        st.markdown(
            """
<div class="glass-card fade-in" style="height: 100%;">
  <h3 style="margin-bottom: 1rem;">🔐 Quản lý phân quyền</h3>
  <p>Trải nghiệm bảo mật và cá nhân hóa với hệ thống vai trò linh hoạt: Admin, Chuyên gia (Expert), Người dân (Citizen) hay Báo chí (Press).</p>
</div>
""",
            unsafe_allow_html=True,
        )
    
    st.markdown("<div style='height: 4rem;'></div>", unsafe_allow_html=True)


def _render_login_page() -> None:
    _render_top_nav()
    st.markdown("<div class='page-title fade-in'>Đăng nhập</div>", unsafe_allow_html=True)
    _render_auth_notice()

    google_ready = bool(config.OAUTH_GOOGLE_CLIENT_ID)
    microsoft_ready = bool(config.OAUTH_MICROSOFT_CLIENT_ID)

    cols = st.columns([1.4, 1])
    with cols[0]:
        st.markdown("#### Đăng nhập bằng email")
        with st.form("email_login_form"):
            email = st.text_input("Email", key="email_login")
            password = st.text_input("Password", type="password", key="email_password")
            submitted = st.form_submit_button("Đăng nhập", use_container_width=True)
        if submitted:
            if not email or not password:
                st.error("Email và mật khẩu là bắt buộc.")
            else:
                user = authenticate_email_user(email, password)
                if user:
                    st.session_state.user = user
                    st.success("Đăng nhập thành công.")
                    st.rerun()
                else:
                    existing = get_user_by_email(email)
                    if (
                        existing
                        and existing.get("provider") == "email"
                        and "email_verified" in existing
                        and not existing.get("email_verified", False)
                    ):
                        st.error("Email chưa được xác thực. Vui lòng kiểm tra email.")
                    else:
                        st.error("Email hoặc mật khẩu không đúng.")

        action_cols = st.columns([1.2, 1.2, 4])
        with action_cols[0]:
            if st.button("Quên mật khẩu", key="login_forgot"):
                _set_page("forgot")
                st.rerun()
        with action_cols[1]:
            if st.button("Đăng ký", key="login_register"):
                _set_page("register")
                st.rerun()

    with cols[1]:
        st.markdown("#### Đăng nhập liên kết")
        if not google_ready:
            st.warning(
                "Google OAuth chưa cấu hình. Cần OAUTH_GOOGLE_CLIENT_ID và "
                "OAUTH_GOOGLE_CLIENT_SECRET."
            )
            st.button("Tiếp tục với Google", use_container_width=True, disabled=True, key="oauth_google_disabled")
        else:
            if "google" not in st.session_state.oauth_states:
                st.session_state.oauth_states["google"] = generate_state("google").split(":", 1)[1]
            google_state = st.session_state.oauth_states["google"]
            auth_url = get_login_url(
                "google",
                config.OAUTH_REDIRECT_URI,
                f"google:{google_state}",
            )
            st.markdown(
                f'<a href="{auth_url}" target="_self" style="display: block; text-align: center; border-radius: 999px; border: 1px solid var(--sea); background-color: white; color: var(--sea-2); padding: 0.45rem 1.2rem; font-weight: 600; text-decoration: none; transition: all 0.2s ease; margin-bottom: 0.5rem;" onmouseover="this.style.backgroundColor=\'var(--sea)\'; this.style.color=\'white\';" onmouseout="this.style.backgroundColor=\'white\'; this.style.color=\'var(--sea-2)\';">Tiếp tục với Google</a>',
                unsafe_allow_html=True
            )

        if not microsoft_ready:
            st.warning(
                "Microsoft OAuth chưa cấu hình. Cần OAUTH_MICROSOFT_CLIENT_ID và "
                "OAUTH_MICROSOFT_CLIENT_SECRET."
            )
            st.button("Tiếp tục với Microsoft", use_container_width=True, disabled=True, key="oauth_microsoft_disabled")
        else:
            if "microsoft" not in st.session_state.oauth_states:
                st.session_state.oauth_states["microsoft"] = generate_state("microsoft").split(":", 1)[1]
            ms_state = st.session_state.oauth_states["microsoft"]
            auth_url = get_login_url(
                "microsoft",
                config.OAUTH_REDIRECT_URI,
                f"microsoft:{ms_state}",
            )
            st.markdown(
                f'<a href="{auth_url}" target="_self" style="display: block; text-align: center; border-radius: 999px; border: 1px solid var(--sea); background-color: white; color: var(--sea-2); padding: 0.45rem 1.2rem; font-weight: 600; text-decoration: none; transition: all 0.2s ease; margin-bottom: 0.5rem;" onmouseover="this.style.backgroundColor=\'var(--sea)\'; this.style.color=\'white\';" onmouseout="this.style.backgroundColor=\'white\'; this.style.color=\'var(--sea-2)\';">Tiếp tục với Microsoft</a>',
                unsafe_allow_html=True
            )


def _render_register_page() -> None:
    _render_top_nav()
    st.markdown("<div class='page-title fade-in'>Đăng ký</div>", unsafe_allow_html=True)
    _render_auth_notice()

    default_role = normalize_role(config.DEFAULT_USER_ROLE)
    if default_role == "admin" or default_role not in NON_ADMIN_ROLES:
        default_role = NON_ADMIN_ROLES[0] if NON_ADMIN_ROLES else "citizen"
    default_index = (
        NON_ADMIN_ROLES.index(default_role) if default_role in NON_ADMIN_ROLES else 0
    )

    with st.form("email_register_form", clear_on_submit=True):
        name = st.text_input("Họ và tên", key="register_name")
        email = st.text_input("Email", key="register_email")
        password = st.text_input("Mật khẩu", type="password", key="register_password")
        confirm = st.text_input("Xác nhận mật khẩu", type="password", key="register_confirm")
        role = st.selectbox(
            "Vai trò",
            options=NON_ADMIN_ROLES,
            format_func=_format_role_label,
            index=default_index,
        )
        submitted = st.form_submit_button("Tạo tài khoản", use_container_width=True)

    if submitted:
        cleaned_name = (name or "").strip()
        cleaned_email = (email or "").strip().lower()
        normalized_role = normalize_role(role)

        if not cleaned_name:
            st.error("Họ và tên là bắt buộc.")
        elif not cleaned_email:
            st.error("Email là bắt buộc.")
        elif not _EMAIL_PATTERN.match(cleaned_email):
            st.error("Vui lòng nhập email hợp lệ.")
        elif len(password or "") < 8:
            st.error("Mật khẩu tối thiểu 8 ký tự.")
        elif password != confirm:
            st.error("Mật khẩu xác nhận không khớp.")
        elif normalized_role == "admin":
            st.error("Không thể chọn vai trò admin khi đăng ký.")
        elif get_user_by_email(cleaned_email):
            st.error("Email đã tồn tại.")
        else:
            user = create_email_user(cleaned_email, cleaned_name, normalized_role, password)
            if user:
                token = create_email_verification_token(cleaned_email)
                if token:
                    verify_link = _build_app_link("verify", token)
                    sent = _send_email(
                        cleaned_email,
                        "Verify your MarineRAG DSS account",
                        f"Verify your email address by visiting this link:\n{verify_link}\n\n"
                        f"This link expires in {config.EMAIL_VERIFY_TOKEN_HOURS} hours.",
                    )
                    if sent:
                        st.success("Tạo tài khoản thành công. Vui lòng kiểm tra email để xác thực.")
                    else:
                        st.info(
                            "Email chưa cấu hình. Dùng link sau để xác thực tài khoản:"
                        )
                        st.code(verify_link)
                else:
                    st.warning("Tạo tài khoản thành công. Hãy yêu cầu gửi lại email xác thực.")
            else:
                st.error("Không thể tạo tài khoản. Vui lòng thử lại.")

    action_cols = st.columns([1.2, 1.2, 4])
    with action_cols[0]:
        if st.button("Đăng nhập", key="register_login"):
            _set_page("login")
            st.rerun()
    with action_cols[1]:
        if st.button("Quên mật khẩu", key="register_forgot"):
            _set_page("forgot")
            st.rerun()

    with st.expander("Gửi lại email xác thực"):
        with st.form("resend_verification"):
            resend_email = st.text_input("Email", key="resend_email")
            resend_submit = st.form_submit_button("Gửi xác thực", use_container_width=True)
        if resend_submit:
            cleaned_email = (resend_email or "").strip().lower()
            if not cleaned_email:
                st.error("Email là bắt buộc.")
            elif not _EMAIL_PATTERN.match(cleaned_email):
                st.error("Vui lòng nhập email hợp lệ.")
            else:
                token = create_email_verification_token(cleaned_email)
                if token:
                    verify_link = _build_app_link("verify", token)
                    sent = _send_email(
                        cleaned_email,
                        "Verify your MarineRAG DSS account",
                        f"Verify your email address by visiting this link:\n{verify_link}\n\n"
                        f"This link expires in {config.EMAIL_VERIFY_TOKEN_HOURS} hours.",
                    )
                    if sent:
                        st.success("Đã gửi email xác thực.")
                    else:
                        st.info("Email chưa cấu hình. Dùng link sau để xác thực:")
                        st.code(verify_link)
                else:
                    st.info("Nếu tài khoản tồn tại và chưa xác thực, bạn sẽ nhận được email.")


def _render_forgot_password_page() -> None:
    _render_top_nav()
    st.markdown("<div class='page-title fade-in'>Quên mật khẩu</div>", unsafe_allow_html=True)
    _render_auth_notice()

    if st.session_state.pending_reset_token:
        st.markdown("#### Đặt lại mật khẩu")
        with st.form("reset_password_form"):
            new_password = st.text_input("Mật khẩu mới", type="password")
            confirm_password = st.text_input("Xác nhận mật khẩu", type="password")
            submitted_reset = st.form_submit_button("Cập nhật mật khẩu", use_container_width=True)
        if submitted_reset:
            if len(new_password or "") < 8:
                st.error("Mật khẩu tối thiểu 8 ký tự.")
            elif new_password != confirm_password:
                st.error("Mật khẩu xác nhận không khớp.")
            elif reset_password_with_token(st.session_state.pending_reset_token, new_password):
                st.session_state.pending_reset_token = ""
                _set_auth_notice("success", "Mật khẩu đã cập nhật. Bạn có thể đăng nhập.")
                _set_page("login")
                st.rerun()
            else:
                st.error("Link đặt lại không hợp lệ hoặc đã hết hạn.")

        st.divider()

    st.markdown("#### Yêu cầu link đặt lại")
    with st.form("password_reset_request"):
        reset_email = st.text_input("Email", key="reset_email")
        reset_submit = st.form_submit_button("Gửi link", use_container_width=True)
    if reset_submit:
        cleaned_email = (reset_email or "").strip().lower()
        if not cleaned_email:
            st.error("Email là bắt buộc.")
        elif not _EMAIL_PATTERN.match(cleaned_email):
            st.error("Vui lòng nhập email hợp lệ.")
        else:
            token = create_password_reset_token(cleaned_email)
            if token:
                reset_link = _build_app_link("reset", token)
                sent = _send_email(
                    cleaned_email,
                    "MarineRAG DSS password reset",
                    f"Use this link to reset your password:\n{reset_link}\n\n"
                    f"This link expires in {config.PASSWORD_RESET_TOKEN_HOURS} hours.",
                )
                if sent:
                    st.success("Đã gửi link đặt lại. Vui lòng kiểm tra email.")
                else:
                    st.info("Email chưa cấu hình. Dùng link sau để đặt lại:")
                    st.code(reset_link)
            else:
                st.info("Nếu tài khoản tồn tại, bạn sẽ nhận được email đặt lại.")

    action_cols = st.columns([1.2, 1.2, 4])
    with action_cols[0]:
        if st.button("Đăng nhập", key="forgot_login"):
            _set_page("login")
            st.rerun()
    with action_cols[1]:
        if st.button("Đăng ký", key="forgot_register"):
            _set_page("register")
            st.rerun()


def _render_public_page() -> None:
    page = st.session_state.page
    if page not in {"home", "login", "register", "forgot"}:
        page = "home"
        st.session_state.page = "home"

    if page == "home":
        _render_home()
    elif page == "login":
        _render_login_page()
    elif page == "register":
        _render_register_page()
    else:
        _render_forgot_password_page()


def _build_source_items(documents: list) -> list[dict]:
    items: dict[str, dict] = {}
    for doc in documents or []:
        metadata = getattr(doc, "metadata", {}) or {}
        source = metadata.get("source")
        if not source:
            continue
        is_external = isinstance(source, str) and source.startswith(("http://", "https://"))
        url = metadata.get("url") or (source if is_external else "")
        title = metadata.get("title") or ""
        entry = items.setdefault(
            source,
            {
                "path": source,
                "url": url,
                "title": title,
                "is_external": is_external,
                "chunk_ids": [],
                "snippet": "",
            },
        )
        chunk_id = metadata.get("chunk_id")
        if chunk_id and chunk_id not in entry["chunk_ids"]:
            entry["chunk_ids"].append(chunk_id)
        if not entry["snippet"]:
            entry["snippet"] = doc.page_content[:400].replace("\n", " ").strip()
    return list(items.values())


def _load_source_preview(file_path: Path, max_chars: int = 8000) -> str | None:
    suffix = file_path.suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv", ".json", ".log", ".xml", ".yaml", ".yml"}:
            content = file_path.read_text(encoding="utf-8", errors="ignore")
            return content[:max_chars] if content else None

        if suffix == ".docx":
            from docx import Document as DocxDocument

            doc = DocxDocument(str(file_path))
            parts: list[str] = [
                paragraph.text.strip()
                for paragraph in doc.paragraphs
                if paragraph.text and paragraph.text.strip()
            ]

            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))

            preview = "\n".join(parts).strip()
            return preview[:max_chars] if preview else None

        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(file_path))
            blocks: list[str] = []
            total_len = 0
            for page in reader.pages[:5]:
                text = (page.extract_text() or "").strip()
                if not text:
                    continue
                blocks.append(text)
                total_len += len(text)
                if total_len >= max_chars:
                    break

            preview = "\n\n".join(blocks).strip()
            return preview[:max_chars] if preview else None

        if suffix in {".xls", ".xlsx"}:
            import pandas as pd

            sheets = pd.read_excel(str(file_path), sheet_name=None, dtype=str)
            sections: list[str] = []
            for sheet_name, frame in list(sheets.items())[:2]:
                if frame is None:
                    continue
                frame = frame.fillna("")
                sections.append(f"[Sheet: {sheet_name}]")
                sections.append(frame.head(20).to_string(index=False))
                if len("\n".join(sections)) >= max_chars:
                    break

            preview = "\n\n".join(sections).strip()
            return preview[:max_chars] if preview else None
    except Exception as exc:
        logger.warning("Failed to preview source '%s': %s", file_path, exc)

    return None


def _render_pdf_inline(
    file_path: Path,
    viewer_key: str,
    max_inline_mb: int = 25,
) -> bool:
    try:
        size_mb = file_path.stat().st_size / (1024 * 1024)
        if size_mb > max_inline_mb:
            st.info(
                f"PDF size {size_mb:.1f} MB is too large for inline view. "
                "Use download to view full file."
            )
            return False

        st.pdf(str(file_path), height=860, key=viewer_key)
        return True
    except Exception as exc:
        logger.warning(
            "Failed to render inline PDF '%s' with native viewer: %s",
            file_path,
            exc,
        )
        return False


def _render_sources(sources, key_prefix: str = "source", expand_items: bool = False) -> None:
    if not sources:
        st.text("No source documents retrieved.")
        return
    if isinstance(sources, str):
        st.text(sources)
        return

    for index, item in enumerate(sources, start=1):
        path = item.get("path", "Unknown")
        if item.get("is_external"):
            label = item.get("title") or path
            st.markdown(f"[{label}]({item.get('url', path)})")
            st.divider()
            continue
        with st.expander(path, expanded=expand_items):
            file_path = Path(path)
            if not file_path.exists():
                st.warning("Source file not found on disk.")
                continue

            suffix = file_path.suffix.lower()
            if suffix == ".pdf":
                rendered = _render_pdf_inline(
                    file_path,
                    viewer_key=f"{key_prefix}_pdf_{index}",
                )
                if not rendered:
                    preview = _load_source_preview(file_path)
                    if preview:
                        st.text_area(
                            "File preview",
                            preview,
                            height=240,
                            key=f"{key_prefix}_preview_{index}",
                        )
                    else:
                        st.info(
                            "Preview not available for this file type. "
                            "Use download to view full file."
                        )
            else:
                preview = _load_source_preview(file_path)
                if preview:
                    st.text_area(
                        "File preview",
                        preview,
                        height=240,
                        key=f"{key_prefix}_preview_{index}",
                    )
                else:
                    st.info(
                        "Preview not available for this file type. "
                        "Use download to view full file."
                    )

            with file_path.open("rb") as handle:
                st.download_button(
                    "Download file",
                    data=handle,
                    file_name=file_path.name,
                    mime="application/octet-stream",
                    key=f"{key_prefix}_download_{index}",
                )


def _render_chat() -> None:
    user = st.session_state.user
    permissions = get_permissions(normalize_role(user.get("role", ""))) if user else {}

    # Display existing messages
    for index, msg in enumerate(st.session_state.messages):
        question_hint = ""
        if msg["role"] == "assistant" and index > 0:
            previous = st.session_state.messages[index - 1]
            if previous.get("role") == "user":
                question_hint = str(previous.get("content", ""))

        with st.chat_message(msg["role"]):
            content = msg["content"]
            if msg["role"] == "assistant" and _is_guide_query(question_hint):
                content = _strip_youtube_urls(content)
            if msg["role"] == "assistant" and not permissions.get("can_view_sources"):
                content = _strip_inline_citations(content)
            st.markdown(content)
            if msg["role"] == "assistant":
                _render_embedded_videos(content, question=question_hint)
            if (
                msg["role"] == "assistant"
                and msg.get("sources")
                and permissions.get("can_view_sources")
            ):
                st.markdown("**📄 Source documents**")
                _render_sources(
                    msg["sources"],
                    key_prefix=f"msg_{index}",
                    expand_items=True,
                )

    # Chat input
    if prompt := st.chat_input(
        "Ask about marine regulations, fish stocks, environmental data…",
        disabled=not st.session_state.setup_done or not user,
    ):
        if not st.session_state.active_session_id:
            title = " ".join(prompt.split()[:6])
            st.session_state.active_session_id = create_chat_session(user["id"], title)

        # Append and display user message
        st.session_state.messages.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)

        # Generate assistant response
        with st.chat_message("assistant"):
            with st.spinner("Searching knowledge base…"):
                try:
                    result = st.session_state.chatbot.chat(prompt)
                    answer = result["answer"]
                    if _is_guide_query(prompt):
                        answer = _strip_youtube_urls(answer)
                    if not permissions.get("can_view_sources"):
                        answer = _strip_inline_citations(answer)
                    source_items = _build_source_items(
                        (result.get("context", []) or [])
                        + (result.get("table_docs", []) or [])
                    )
                except Exception as exc:
                    logger.exception("Error during chat.")
                    answer = f"⚠️ An error occurred: {exc}"
                    source_items = []

            st.markdown(answer)
            _render_embedded_videos(answer, question=prompt)
            if source_items and permissions.get("can_view_sources"):
                st.markdown("**📄 Source documents**")
                _render_sources(
                    source_items,
                    key_prefix=f"live_{len(st.session_state.messages)}",
                    expand_items=True,
                )

        st.session_state.messages.append(
            {"role": "assistant", "content": answer, "sources": source_items}
        )

        append_message(
            st.session_state.active_session_id,
            user["id"],
            "user",
            prompt,
        )
        append_message(
            st.session_state.active_session_id,
            user["id"],
            "assistant",
            answer,
            sources=source_items,
        )


def _render_example_queries() -> None:
    user = st.session_state.user
    permissions = (
        get_permissions(normalize_role(user.get("role", "")))
        if user
        else get_permissions("citizen")
    )
    if st.session_state.setup_done and not st.session_state.messages:
        st.markdown("### 💡 Example questions")
        examples = [
            "What is the minimum conservation reference size for Atlantic Cod?",
            "What are the main objectives of the EU Common Fisheries Policy?",
            "What temperature range is acceptable for healthy marine ecosystems?",
            "Explain the difference between MSY and MEY in fisheries management.",
            "What are the regulations under MARPOL regarding plastic discharge?",
            "What is the 30×30 marine protection target and when must it be achieved?",
            "How are Bluefin tuna stocks currently classified?",
            "What fishing gear is prohibited near deep-sea sensitive areas?",
        ]
        cols = st.columns(2)
        for i, example in enumerate(examples):
            if cols[i % 2].button(example, use_container_width=True, key=f"ex_{i}"):
                # Trigger a chat with the example question
                st.session_state.messages.append({"role": "user", "content": example})
                with st.spinner("Searching knowledge base…"):
                    try:
                        result = st.session_state.chatbot.chat(example)
                        answer = result["answer"]
                        if _is_guide_query(example):
                            answer = _strip_youtube_urls(answer)
                        if not permissions.get("can_view_sources"):
                            answer = _strip_inline_citations(answer)
                        source_items = _build_source_items(
                            (result.get("context", []) or [])
                            + (result.get("table_docs", []) or [])
                        )
                    except Exception as exc:
                        logger.exception("Error during example chat.")
                        answer = f"⚠️ An error occurred: {exc}"
                        source_items = []
                st.session_state.messages.append(
                    {"role": "assistant", "content": answer, "sources": source_items}
                )
                st.rerun()


def _render_history() -> None:
    user = st.session_state.user
    if not user:
        st.info("Please sign in to view history.")
        return

    permissions = get_permissions(normalize_role(user.get("role", "")))
    if permissions.get("can_view_all_sessions"):
        users = list_users()
        user_lookup = {u.get("id"): u for u in users}
        user_options = ["All"] + [u.get("email", u.get("id", "")) for u in users]
        selection = st.selectbox("Filter by user", user_options)
        filter_user_id = None
        if selection != "All":
            for u in users:
                if u.get("email") == selection:
                    filter_user_id = u.get("id")
                    break
        sessions = list_sessions(filter_user_id)
    else:
        user_lookup = {user.get("id"): user}
        sessions = list_sessions(user.get("id"))

    if not sessions:
        st.info("No chat sessions found.")
        return

    labels = [f"{s.get('title','(no title)')} ({s.get('updated_at')})" for s in sessions]
    index = st.selectbox("Chat sessions", list(range(len(sessions))), format_func=lambda i: labels[i])
    session = sessions[index]
    if st.button("Load session", use_container_width=True):
        st.session_state.active_session_id = session["id"]
        messages = get_session_messages(session["id"])
        st.session_state.messages = [
            {
                "role": msg.get("role"),
                "content": msg.get("content"),
                "sources": msg.get("sources"),
            }
            for msg in messages
        ]
        st.rerun()

    if permissions.get("can_view_all_sessions"):
        owner = user_lookup.get(session.get("user_id"), {})
        st.caption(f"Owner: {owner.get('email', session.get('user_id'))}")


def _render_admin() -> None:
    user = st.session_state.user
    if not user:
        return
    permissions = get_permissions(normalize_role(user.get("role", "")))
    if not permissions.get("can_view_admin"):
        st.info("Admin access required.")
        return

    stats = get_admin_stats()
    st.metric("Users", stats.get("users", 0))
    st.metric("Sessions", stats.get("sessions", 0))
    st.metric("Messages", stats.get("messages", 0))

    st.subheader("Filters")
    users = list_users()
    role_options = ["All"] + ROLES
    role_filter = st.selectbox("Role filter", role_options, index=0)

    filtered_users = [u for u in users if role_filter == "All" or u.get("role") == role_filter]
    user_options = ["All"] + [u.get("email", u.get("id", "")) for u in filtered_users]
    user_selection = st.selectbox("User filter", user_options, index=0)

    date_range = st.date_input(
        "Date range",
        value=(datetime.now().date(), datetime.now().date()),
    )
    if isinstance(date_range, tuple):
        if len(date_range) == 2:
            start_date, end_date = date_range
        elif len(date_range) == 1:
            start_date = end_date = date_range[0]
        else:
            start_date = end_date = datetime.now().date()
    else:
        start_date = end_date = date_range

    start_dt = datetime.combine(start_date, time.min, tzinfo=timezone.utc)
    end_dt = datetime.combine(end_date, time.max, tzinfo=timezone.utc)

    selected_user_ids = []
    if user_selection != "All":
        for entry in filtered_users:
            if entry.get("email") == user_selection:
                selected_user_ids = [entry.get("id")]
                break
    elif role_filter != "All":
        selected_user_ids = [str(value) for value in list_user_ids_by_role(role_filter)]

    user_ids = [
        ObjectId(uid) for uid in selected_user_ids
    ] if selected_user_ids else None

    st.subheader("Analytics")
    analytics = get_admin_analytics(user_ids=user_ids, start=start_dt, end=end_dt)
    st.metric("Messages (filtered)", analytics.get("messages_total", 0))
    st.metric("Sessions (filtered)", analytics.get("sessions_total", 0))
    st.metric("Active users", analytics.get("active_users", 0))

    st.markdown("**Messages per day**")
    st.dataframe(analytics.get("messages_per_day", []), use_container_width=True)

    st.markdown("**Sessions per day**")
    st.dataframe(analytics.get("sessions_per_day", []), use_container_width=True)

    st.markdown("**Top users by message count**")
    st.dataframe(analytics.get("top_users", []), use_container_width=True)

    st.markdown("**Role breakdown**")
    st.dataframe(analytics.get("role_breakdown", []), use_container_width=True)

    st.subheader("Data exports")

    def _export_csv(rows: list[dict], fieldnames: list[str]) -> bytes:
        buffer = io.StringIO()
        writer = csv.DictWriter(buffer, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({name: row.get(name, "") for name in fieldnames})
        return buffer.getvalue().encode("utf-8")

    sessions = list_sessions_filtered(user_ids=user_ids, start=start_dt, end=end_dt, limit=5000)
    session_rows = [
        {
            "id": s.get("id"),
            "title": s.get("title"),
            "user_id": s.get("user_id"),
            "created_at": s.get("created_at"),
            "updated_at": s.get("updated_at"),
        }
        for s in sessions
    ]
    st.download_button(
        "Export sessions (CSV)",
        data=_export_csv(session_rows, ["id", "title", "user_id", "created_at", "updated_at"]),
        file_name="chat_sessions.csv",
        mime="text/csv",
        key="export_sessions",
    )
    st.dataframe(session_rows, use_container_width=True)

    message_role_filter = st.selectbox("Message role filter", ["All", "user", "assistant"], index=0)
    message_role = None if message_role_filter == "All" else message_role_filter
    messages = list_messages_filtered(
        user_ids=user_ids,
        start=start_dt,
        end=end_dt,
        role=message_role,
        limit=5000,
    )
    message_rows = [
        {
            "id": m.get("id"),
            "session_id": m.get("session_id"),
            "user_id": m.get("user_id"),
            "role": m.get("role"),
            "content": m.get("content"),
            "created_at": m.get("created_at"),
        }
        for m in messages
    ]
    st.download_button(
        "Export messages (CSV)",
        data=_export_csv(
            message_rows,
            ["id", "session_id", "user_id", "role", "content", "created_at"],
        ),
        file_name="chat_messages.csv",
        mime="text/csv",
        key="export_messages",
    )
    st.dataframe(message_rows, use_container_width=True)

    st.subheader("User roles")
    for entry in users:
        cols = st.columns([3, 3, 2])
        cols[0].write(entry.get("email", ""))
        cols[1].write(entry.get("name", ""))
        current_role = entry.get("role", "citizen")
        new_role = cols[2].selectbox(
            "Role",
            ROLES,
            index=ROLES.index(current_role) if current_role in ROLES else 0,
            key=f"role_{entry.get('id')}",
        )
        if new_role != current_role:
            update_user_role(entry["id"], new_role)
            st.rerun()


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main() -> None:
    _init_session_state()
    ensure_indexes()
    _handle_oauth_callback()
    if not st.session_state.user:
        _handle_auth_query_params()
        _inject_public_styles()
        _render_public_page()
        return

    _render_sidebar()
    _render_header()

    view_options = ["Chat", "History"]
    permissions = get_permissions(normalize_role(st.session_state.user.get("role", "")))
    if permissions.get("can_view_admin"):
        view_options.append("Admin")
    view = st.sidebar.radio("View", view_options)

    if not st.session_state.setup_done:
        st.info(
            "👈 Select an LLM mode in the sidebar and click **Initialise** to start."
        )
        return

    if view == "Chat":
        _render_example_queries()
        _render_chat()
    elif view == "History":
        _render_history()
    else:
        _render_admin()


if __name__ == "__main__":
    main()
